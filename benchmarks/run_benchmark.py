from __future__ import annotations

import argparse
import csv
import json
from collections import defaultdict
from pathlib import Path

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from benchmark_corpus import BENCHMARK_VERSION, BenchmarkCase, make_benchmark_corpus, validate_corpus
from circular_interpolation_inertia import (
    CircularInterpolationResult,
    interpolate_circular_inertial,
    interpolate_circular_linear,
)
from motor_profile_evaluation import (
    interpolate_circular_boundary_hermite,
    interpolate_circular_constant_acceleration,
)


MODELS = (
    "numpy_unwrap_linear",
    "constant_velocity_rts",
    "constant_acceleration_rts",
    "boundary_velocity_hermite",
)


def insert_held_gap(values: np.ndarray, start: int, missing_samples: int) -> tuple[np.ndarray, np.ndarray, int]:
    observed = np.asarray(values, dtype=float).copy()
    invalid = np.zeros(observed.size, dtype=bool)
    end = min(observed.size, start + missing_samples)
    if start < 1 or end >= observed.size:
        raise ValueError("Gap must have a valid sample on both sides")
    observed[start:end] = observed[start - 1]
    invalid[start:end] = True
    return observed, invalid, end


def align_to_truth(estimate: np.ndarray, truth: np.ndarray, anchor: int) -> np.ndarray:
    offset = round((truth[anchor] - estimate[anchor]) / 360.0) * 360.0
    return estimate + offset


def model_outputs(time_s: np.ndarray, observed: np.ndarray, invalid: np.ndarray) -> dict[str, np.ndarray]:
    _, linear_u = interpolate_circular_linear(time_s, observed, invalid)
    cv: CircularInterpolationResult = interpolate_circular_inertial(
        time_s,
        observed,
        invalid,
        velocity_random_walk_std=50_000.0,
        measurement_std_deg=0.03,
        gate_sigma=8.0,
    )
    ca = interpolate_circular_constant_acceleration(
        time_s,
        observed,
        invalid,
        acceleration_random_walk_std=2_000_000.0,
        measurement_std_deg=0.03,
        gate_sigma=8.0,
    )
    hermite = interpolate_circular_boundary_hermite(time_s, observed, invalid)
    return {
        "numpy_unwrap_linear": linear_u,
        "constant_velocity_rts": cv.unwrapped_deg,
        "constant_acceleration_rts": ca.unwrapped_deg,
        "boundary_velocity_hermite": hermite.unwrapped_deg,
    }


def evaluate_case_window(
    case: BenchmarkCase,
    *,
    window_label: str,
    center_s: float,
    motion_hint: str,
    missing_samples: int,
    local_margin_ms: float = 250.0,
) -> list[dict[str, object]]:
    fs = case.sample_rate_hz
    center_index = int(round(center_s * fs))
    start = center_index - missing_samples // 2
    start = max(1, min(start, case.time_s.size - missing_samples - 1))

    observed, invalid, end = insert_held_gap(
        case.measured_angle_deg, start, missing_samples
    )

    margin_samples = int(round(max(local_margin_ms, 1.5 * missing_samples * 1000.0 / fs + 80.0) * fs / 1000.0))
    lo = max(0, start - margin_samples)
    hi = min(case.time_s.size, end + margin_samples + 1)

    local_t = case.time_s[lo:hi]
    local_observed = observed[lo:hi]
    local_invalid = invalid[lo:hi]
    local_truth = case.true_angle_unwrapped_deg[lo:hi]
    local_velocity = case.true_velocity_deg_s[lo:hi]
    local_acceleration = case.true_acceleration_deg_s2[lo:hi]

    local_start = start - lo
    local_end = end - lo
    anchor = local_start - 1
    check_index = local_end

    outputs = model_outputs(local_t, local_observed, local_invalid)
    rows: list[dict[str, object]] = []

    true_gap_velocity = local_velocity[local_invalid]
    true_gap_acceleration = local_acceleration[local_invalid]
    endpoint_span_ms = (missing_samples + 1) * 1000.0 / fs
    true_gap_displacement_deg = float(local_truth[local_end] - local_truth[anchor])

    for model, estimate in outputs.items():
        aligned = align_to_truth(estimate, local_truth, anchor)
        errors = aligned[local_invalid] - local_truth[local_invalid]
        endpoint_error = float(aligned[check_index] - local_truth[check_index])
        rows.append(
            {
                "benchmark_version": BENCHMARK_VERSION,
                "case_id": case.case_id,
                "family": case.family,
                "description": case.description,
                "window_label": window_label,
                "motion_hint": motion_hint,
                "center_s": center_s,
                "sample_rate_hz": fs,
                "missing_samples": missing_samples,
                "held_duration_ms": missing_samples * 1000.0 / fs,
                "endpoint_span_ms": endpoint_span_ms,
                "model": model,
                "true_start_speed_rpm": float(local_velocity[anchor] / 6.0),
                "true_end_speed_rpm": float(local_velocity[check_index] / 6.0),
                "gap_max_abs_speed_rpm": float(np.max(np.abs(true_gap_velocity)) / 6.0),
                "gap_velocity_span_rpm": float((np.max(true_gap_velocity) - np.min(true_gap_velocity)) / 6.0),
                "gap_max_abs_acceleration_deg_s2": float(np.max(np.abs(true_gap_acceleration))),
                "true_endpoint_displacement_deg": true_gap_displacement_deg,
                "unwrapped_rmse_deg": float(np.sqrt(np.mean(errors**2))),
                "mean_abs_error_deg": float(np.mean(np.abs(errors))),
                "max_abs_error_deg": float(np.max(np.abs(errors))),
                "endpoint_error_deg": endpoint_error,
                "endpoint_abs_error_deg": abs(endpoint_error),
                "winding_slip": abs(endpoint_error) > 180.0,
            }
        )
    return rows


def run_evaluation(
    cases: list[BenchmarkCase],
    missing_sample_counts: tuple[int, ...],
) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    total_windows = sum(len(case.gap_windows) for case in cases)
    completed = 0
    for case in cases:
        for window in case.gap_windows:
            for missing_samples in missing_sample_counts:
                rows.extend(
                    evaluate_case_window(
                        case,
                        window_label=window.label,
                        center_s=window.center_s,
                        motion_hint=window.motion_hint,
                        missing_samples=missing_samples,
                    )
                )
            completed += 1
            if completed % 20 == 0 or completed == total_windows:
                print(f"Evaluated {completed}/{total_windows} fixed windows", flush=True)
    return rows


def summarize_results(rows: list[dict[str, object]]) -> tuple[list[dict[str, object]], list[dict[str, object]]]:
    # First aggregate windows within each case, so profiles with more event windows
    # do not receive a larger weight in the global result.
    per_case_buckets: dict[tuple[str, int, str], list[dict[str, object]]] = defaultdict(list)
    for row in rows:
        per_case_buckets[(str(row["case_id"]), int(row["missing_samples"]), str(row["model"]))].append(row)

    per_case: list[dict[str, object]] = []
    for (case_id, missing_samples, model), bucket in per_case_buckets.items():
        rmse = np.asarray([float(r["unwrapped_rmse_deg"]) for r in bucket])
        slips = np.asarray([bool(r["winding_slip"]) for r in bucket])
        per_case.append(
            {
                "benchmark_version": BENCHMARK_VERSION,
                "case_id": case_id,
                "family": bucket[0]["family"],
                "missing_samples": missing_samples,
                "held_duration_ms": bucket[0]["held_duration_ms"],
                "endpoint_span_ms": bucket[0]["endpoint_span_ms"],
                "model": model,
                "window_count": len(bucket),
                "case_median_rmse_deg": float(np.median(rmse)),
                "case_p95_rmse_deg": float(np.percentile(rmse, 95)),
                "case_max_rmse_deg": float(np.max(rmse)),
                "case_winding_slip_percent": 100.0 * float(np.mean(slips)),
                "case_any_winding_slip": bool(np.any(slips)),
            }
        )

    global_buckets: dict[tuple[int, str], list[dict[str, object]]] = defaultdict(list)
    for row in per_case:
        global_buckets[(int(row["missing_samples"]), str(row["model"]))].append(row)

    summary: list[dict[str, object]] = []
    for (missing_samples, model), bucket in global_buckets.items():
        medians = np.asarray([float(r["case_median_rmse_deg"]) for r in bucket])
        p95s = np.asarray([float(r["case_p95_rmse_deg"]) for r in bucket])
        any_slips = np.asarray([bool(r["case_any_winding_slip"]) for r in bucket])
        all_window_rows = [
            row for row in rows
            if int(row["missing_samples"]) == missing_samples and str(row["model"]) == model
        ]
        window_rmse = np.asarray([float(r["unwrapped_rmse_deg"]) for r in all_window_rows])
        window_slips = np.asarray([bool(r["winding_slip"]) for r in all_window_rows])
        summary.append(
            {
                "benchmark_version": BENCHMARK_VERSION,
                "missing_samples": missing_samples,
                "held_duration_ms": bucket[0]["held_duration_ms"],
                "endpoint_span_ms": bucket[0]["endpoint_span_ms"],
                "model": model,
                "case_count": len(bucket),
                "window_count": len(all_window_rows),
                "balanced_median_rmse_deg": float(np.median(medians)),
                "balanced_p90_case_median_rmse_deg": float(np.percentile(medians, 90)),
                "balanced_p95_case_median_rmse_deg": float(np.percentile(medians, 95)),
                "balanced_p95_case_p95_rmse_deg": float(np.percentile(p95s, 95)),
                "worst_case_median_rmse_deg": float(np.max(medians)),
                "window_winding_slip_percent": 100.0 * float(np.mean(window_slips)),
                "cases_with_any_slip_percent": 100.0 * float(np.mean(any_slips)),
                "windows_below_0_1deg_percent": 100.0 * float(np.mean(window_rmse < 0.1)),
                "windows_below_1deg_percent": 100.0 * float(np.mean(window_rmse < 1.0)),
                "windows_below_5deg_percent": 100.0 * float(np.mean(window_rmse < 5.0)),
                "windows_below_30deg_percent": 100.0 * float(np.mean(window_rmse < 30.0)),
            }
        )

    per_case.sort(key=lambda r: (int(r["missing_samples"]), str(r["case_id"]), str(r["model"])))
    summary.sort(key=lambda r: (int(r["missing_samples"]), MODELS.index(str(r["model"]))))
    return per_case, summary


def write_csv(rows: list[dict[str, object]], path: Path) -> None:
    if not rows:
        return
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)


def make_rmse_matrix(rows: list[dict[str, object]]) -> list[dict[str, object]]:
    lookup = {
        (str(r["case_id"]), str(r["window_label"]), int(r["missing_samples"]), str(r["model"])): float(r["unwrapped_rmse_deg"])
        for r in rows
    }
    identities = sorted({(str(r["case_id"]), str(r["family"]), str(r["window_label"]), str(r["motion_hint"])) for r in rows})
    gaps = sorted({int(r["missing_samples"]) for r in rows})
    matrix: list[dict[str, object]] = []
    for case_id, family, window_label, motion_hint in identities:
        row: dict[str, object] = {
            "case_id": case_id,
            "family": family,
            "window_label": window_label,
            "motion_hint": motion_hint,
        }
        for model in MODELS:
            for gap in gaps:
                row[f"{model}__missing_{gap}"] = lookup[(case_id, window_label, gap, model)]
        matrix.append(row)
    return matrix


def summarize_numpy_by_speed_tier(
    rows: list[dict[str, object]],
    cases: list[BenchmarkCase],
) -> list[dict[str, object]]:
    case_speed = {case.case_id: case.max_abs_speed_rpm for case in cases}
    limits = (500, 1500, 3000, 6000, 10000)
    output: list[dict[str, object]] = []
    linear_rows = [r for r in rows if str(r["model"]) == "numpy_unwrap_linear"]
    for limit in limits:
        for missing_samples in sorted({int(r["missing_samples"]) for r in linear_rows}):
            bucket = [
                r for r in linear_rows
                if int(r["missing_samples"]) == missing_samples
                and case_speed[str(r["case_id"])] <= limit + 1e-9
            ]
            rmse = np.asarray([float(r["unwrapped_rmse_deg"]) for r in bucket])
            slips = np.asarray([bool(r["winding_slip"]) for r in bucket])
            output.append(
                {
                    "benchmark_version": BENCHMARK_VERSION,
                    "max_dataset_speed_rpm": limit,
                    "missing_samples": missing_samples,
                    "held_duration_ms": bucket[0]["held_duration_ms"],
                    "endpoint_span_ms": bucket[0]["endpoint_span_ms"],
                    "dataset_count": len({str(r["case_id"]) for r in bucket}),
                    "window_count": len(bucket),
                    "median_unwrapped_rmse_deg": float(np.median(rmse)),
                    "p95_unwrapped_rmse_deg": float(np.percentile(rmse, 95)),
                    "winding_slip_percent": 100.0 * float(np.mean(slips)),
                    "windows_below_1deg_percent": 100.0 * float(np.mean(rmse < 1.0)),
                }
            )
    return output


def create_workbook(
    rows: list[dict[str, object]],
    summary: list[dict[str, object]],
    numpy_speed_tiers: list[dict[str, object]],
    cases: list[BenchmarkCase],
    output_path: Path,
) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    dark_fill = PatternFill("solid", fgColor="1F4E78")
    sub_fill = PatternFill("solid", fgColor="D9EAF7")
    white_font = Font(color="FFFFFF", bold=True)
    header_font = Font(bold=True)

    ws = wb.create_sheet("Overview")
    overview = [
        ("Benchmark version", BENCHMARK_VERSION),
        ("Fixed datasets", len(cases)),
        ("Fixed event windows", sum(len(c.gap_windows) for c in cases)),
        ("Models", len(MODELS)),
        ("Primary error", "Unwrapped RMSE in the missing samples"),
        ("Look-ahead", "Enabled for all methods; samples after the gap are used"),
        ("Weighting", "Each dataset has equal weight in global summaries"),
        ("Gap convention", "missing_samples=N; valid boundary samples are N+1 sample intervals apart"),
    ]
    ws.append(["Motor angle interpolation benchmark", ""])
    ws.merge_cells("A1:B1")
    ws["A1"].fill = dark_fill
    ws["A1"].font = white_font
    ws["A1"].alignment = Alignment(horizontal="center")
    for item in overview:
        ws.append(list(item))
    ws.column_dimensions["A"].width = 27
    ws.column_dimensions["B"].width = 95

    # Summary sheet.
    ws = wb.create_sheet("Balanced summary")
    headers = list(summary[0].keys())
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = dark_fill
        cell.font = white_font
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for item in summary:
        ws.append([item[h] for h in headers])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col, header in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(col)].width = max(13, min(34, len(header) + 2))
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = "0.000"

    ws = wb.create_sheet("NumPy by speed")
    headers = list(numpy_speed_tiers[0].keys())
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = dark_fill
        cell.font = white_font
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for item in numpy_speed_tiers:
        ws.append([item[h] for h in headers])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col, header in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(col)].width = max(13, min(31, len(header) + 2))
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = "0.000"

    gaps = sorted({int(r["missing_samples"]) for r in rows})
    for model in MODELS:
        ws = wb.create_sheet(model[:31])
        model_rows = [r for r in rows if str(r["model"]) == model]
        lookup = {
            (str(r["case_id"]), str(r["window_label"]), int(r["missing_samples"])): float(r["unwrapped_rmse_deg"])
            for r in model_rows
        }
        identities = sorted({(str(r["case_id"]), str(r["family"]), str(r["window_label"]), str(r["motion_hint"])) for r in model_rows})
        headers = ["case_id", "family", "window_label", "motion_hint"] + [f"{g} missing" for g in gaps]
        ws.append(headers)
        for cell in ws[1]:
            cell.fill = dark_fill
            cell.font = white_font
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
        for case_id, family, window_label, motion_hint in identities:
            ws.append(
                [case_id, family, window_label, motion_hint]
                + [lookup[(case_id, window_label, gap)] for gap in gaps]
            )
        ws.freeze_panes = "E2"
        ws.auto_filter.ref = ws.dimensions
        for col in range(1, 5):
            ws.column_dimensions[get_column_letter(col)].width = (29, 20, 27, 22)[col - 1]
        for col in range(5, 5 + len(gaps)):
            ws.column_dimensions[get_column_letter(col)].width = 13
            for cell in ws[get_column_letter(col)][1:]:
                cell.number_format = "0.000"
        data_range = f"E2:{get_column_letter(4 + len(gaps))}{ws.max_row}"
        ws.conditional_formatting.add(
            data_range,
            ColorScaleRule(
                start_type="num", start_value=0, start_color="63BE7B",
                mid_type="num", mid_value=5, mid_color="FFEB84",
                end_type="num", end_value=180, end_color="F8696B",
            ),
        )

    ws = wb.create_sheet("Dataset manifest")
    manifest_headers = [
        "case_id", "family", "description", "duration_s", "sample_count", "seed",
        "max_abs_speed_rpm", "max_abs_acceleration_deg_s2", "window_count", "metadata_json"
    ]
    ws.append(manifest_headers)
    for cell in ws[1]:
        cell.fill = dark_fill
        cell.font = white_font
    for case in cases:
        ws.append([
            case.case_id,
            case.family,
            case.description,
            float(case.time_s[-1]),
            int(case.time_s.size),
            case.seed,
            case.max_abs_speed_rpm,
            case.max_abs_acceleration_deg_s2,
            len(case.gap_windows),
            json.dumps(case.metadata, sort_keys=True),
        ])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = [30, 20, 64, 13, 14, 11, 20, 29, 14, 45]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    for row in ws.iter_rows(min_row=2):
        row[2].alignment = Alignment(wrap_text=True, vertical="top")
        row[9].alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(output_path)
    # Structural validation: ensure the file can be loaded and expected sheets exist.
    checked = load_workbook(output_path, read_only=True, data_only=False)
    expected = {"Overview", "Balanced summary", "NumPy by speed", "Dataset manifest", *[m[:31] for m in MODELS]}
    if not expected.issubset(set(checked.sheetnames)):
        raise AssertionError("Workbook validation failed")
    checked.close()


def print_short_summary(summary: list[dict[str, object]]) -> None:
    print("\nBalanced corpus result (each dataset has equal weight)")
    for gap in sorted({int(r["missing_samples"]) for r in summary}):
        candidates = [r for r in summary if int(r["missing_samples"]) == gap]
        best = min(
            candidates,
            key=lambda r: (
                float(r["cases_with_any_slip_percent"]),
                float(r["balanced_p95_case_p95_rmse_deg"]),
                float(r["balanced_median_rmse_deg"]),
            ),
        )
        linear = next(r for r in candidates if r["model"] == "numpy_unwrap_linear")
        print(
            f"{gap:3d} missing: best={str(best['model']):30s} "
            f"median={float(best['balanced_median_rmse_deg']):8.3f} deg; "
            f"NumPy median={float(linear['balanced_median_rmse_deg']):8.3f} deg, "
            f"slip windows={float(linear['window_winding_slip_percent']):5.1f}%"
        )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-dir", type=Path, default=Path(__file__).resolve().parent)
    parser.add_argument(
        "--missing-samples",
        type=int,
        nargs="+",
        default=[1, 2, 3, 4, 5, 10, 20, 50, 100, 200],
    )
    args = parser.parse_args()

    args.output_dir.mkdir(parents=True, exist_ok=True)
    cases = make_benchmark_corpus()
    validate_corpus(cases)
    rows = run_evaluation(cases, tuple(args.missing_samples))
    per_case, summary = summarize_results(rows)
    matrix = make_rmse_matrix(rows)
    numpy_speed_tiers = summarize_numpy_by_speed_tier(rows, cases)

    write_csv(rows, args.output_dir / "benchmark_window_results.csv")
    write_csv(per_case, args.output_dir / "benchmark_case_balanced_results.csv")
    write_csv(summary, args.output_dir / "benchmark_balanced_summary.csv")
    write_csv(matrix, args.output_dir / "benchmark_rmse_matrix.csv")
    write_csv(numpy_speed_tiers, args.output_dir / "benchmark_numpy_speed_tiers.csv")
    create_workbook(
        rows,
        summary,
        numpy_speed_tiers,
        cases,
        args.output_dir / "motor_interpolation_benchmark_matrix.xlsx",
    )
    print_short_summary(summary)


if __name__ == "__main__":
    main()
